"""Generates docs/product/Hydris-Production-Readiness-Plan.xlsx (inventory, corrections, enhancements, 30-day plan)."""
import datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT = Path(__file__).resolve().parents[1] / 'Hydris-Production-Readiness-Plan.xlsx'
NAVY, BLUE, TINT, BORDER = '0C1E85', '1954F1', 'EEF2FE', 'D5DEF7'
thin = Side(style='thin', color=BORDER)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# ------------------------------------------------------------------ inventory
BACKEND = [
    # area, component, what it does, status, location
    ('API modules', 'Auth (17 routes)', 'Password login, sessions with rotating refresh, MFA (TOTP), password reset, invitations, demo login with named guest accounts', 'Built', 'apps/api/src/modules/auth.ts'),
    ('API modules', 'Organization (13)', 'Org settings, members, roles, invitations, audit log, platform admin', 'Built', 'modules/org.ts'),
    ('API modules', 'Plant (28)', 'Sites, process areas, assets, sensors, plant map, documents/SOPs, attachments', 'Built', 'modules/plant.ts'),
    ('API modules', 'Telemetry (6)', 'HTTPS ingest with API keys, idempotent readings, history, corrections, freshness', 'Built', 'modules/telemetry.ts'),
    ('API modules', 'Alarms (9)', 'Rule evaluation, lifecycle (acknowledge, investigate, resolve with note, close), assignment', 'Built', 'modules/alarms.ts'),
    ('API modules', 'Inspections (15)', 'Templates, rounds, step responses with offline op replay, submit, review', 'Built', 'modules/inspections.ts'),
    ('API modules', 'Maintenance (7)', 'Work requests, assignment, execution, activity log', 'Built', 'modules/maintenance.ts'),
    ('API modules', 'Devices and pairing (16)', 'RFC 8628-style pairing (300 s code), device tokens, suspend/revoke/reassign, glasses logout', 'Built', 'modules/devices.ts'),
    ('API modules', 'UI profiles (13)', 'What each glasses screen shows; draft, publish, rollback, live push', 'Built', 'modules/ui-profiles.ts'),
    ('API modules', 'Audit documents (12)', 'Scan upload, OCR structuring, cross-check, plausibility, vision enhance, field review, second-person approval, searchable PDF', 'Built', 'modules/audit-docs.ts, core/audit-pdf.ts, core/vision-extract.ts'),
    ('API modules', 'Voice assistant (1)', 'Grounded answers from account data: readings, alarms, my work, assets, SOPs; Claude tool use or rules engine', 'Built', 'modules/voice-assistant.ts'),
    ('API modules', 'Insights (4)', 'Overview dashboard, reports, search, assist', 'Built', 'modules/insights.ts'),
    ('API modules', 'Realtime (1)', 'Server-sent events (18 domain event types), memory or db mode', 'Built', 'modules/realtime.ts, app/api/events/stream'),
    ('API modules', 'Simulation (3) + Public (1)', 'Demo scenarios and synthetic telemetry; contact / demo request', 'Built', 'modules/simulation.ts, public.ts'),
    ('Platform', 'Health and cron', '/api/health, /api/cron (daily on Vercel)', 'Partial', 'vercel.json, jobs.ts'),
    ('Platform', 'Tenancy and RLS', 'Postgres row-level security; withTenant/withSystem; 36 of 42 tables with RLS', 'Built', 'core/*, drizzle/*.sql'),
    ('Platform', 'Demo mode', 'In-memory per instance, shared state via Postgres kv / Vercel cache, stateless demo device tokens', 'Built', 'core/demo-shared.ts, core/demo-guests.ts'),
    ('Data', 'Database schema', '42 tables, 548 columns, 23 enums, 9 migrations, 71 indexes', 'Built', 'apps/api/src/db/schema.ts'),
    ('Domain', 'Domain package', 'Audit doc structuring, plausibility, voice wake/measurand detection, roles (38 permissions, 6 roles), UI profile rules', 'Built', 'packages/domain/src'),
    ('Quality', 'Automated tests', '67 passing: 26 domain + 41 API; 4 Playwright journeys', 'Built', 'packages/domain/test, apps/api/test, tests/e2e'),
    ('Deployment', 'Vercel + Postgres', 'Next.js with API in-process at hydrisview.vercel.app; Neon-ready', 'Built', 'vercel.json, apps/web/src/app/api/[...path]'),
    ('Email', 'Mail delivery', 'Messages only written to mail_outbox; nothing is sent', 'Partial', 'mail_outbox'),
    ('Storage', 'File storage', 'Audit page images stored base64 in Postgres', 'Partial', 'audit_document_pages'),
    ('CI/CD', 'Pipeline', 'No GitHub Actions; deploys straight from main', 'Planned', '.github (missing)'),
]
FRONTEND = [
    ('Website', 'Public site', 'Home, product, technology, operations, demo, contact, privacy; demo request', 'Built', 'apps/web/src/app/(site)'),
    ('Website', 'Login and demo', 'Sign in, MFA, forgot/reset password, one-click named demo account', 'Built', 'app/login'),
    ('Console', 'Overview', 'KPIs, process snapshot, alarm frequency, connectivity', 'Built', 'app/console/page.tsx'),
    ('Console', 'Alarms', 'List, severity filters, lifecycle actions, notes', 'Built', 'console/alarms'),
    ('Console', 'Telemetry', 'Sensor cards and table, history charts, freshness', 'Built', 'console/telemetry'),
    ('Console', 'Plant: sites, assets, map', 'Hierarchy management and plant map with routes', 'Built', 'console/sites, assets, map'),
    ('Console', 'Inspections and maintenance', 'Templates, pipeline, review, work orders', 'Built', 'console/inspections, maintenance'),
    ('Console', 'Documents / SOPs', 'Approved procedures, versions, search', 'Built', 'console/documents'),
    ('Console', 'Field glasses and profiles', 'Pairing approval, device states, UI profile editor', 'Built', 'console/devices, glasses-profiles'),
    ('Console', 'Audit documents', 'Review queue, field-by-field verify, approve, PDF download', 'Built', 'console/audit-docs'),
    ('Console', 'Reports, operators, audit log, settings, platform', 'Admin and reporting screens', 'Built', 'console/reports, operators, audit, settings, platform'),
    ('Glasses web app', 'HUD and phone modes', '12 screens, back navigation, menu, logout/unpair, offline inspection queue, voice, audit scan with guided capture', 'Built', 'apps/web/src/glasses'),
    ('Glasses web app', 'Simulator and pairing page', 'Browser glasses simulator with scenarios; /pair approval', 'Built', 'app/simulator, app/pair'),
    ('Android', 'Hydris Field app (Kotlin/Compose)', 'Login + demo, Home dashboard, Work (alarms/inspections/tasks/sensors), Scan (glasses, phone, import), Hydris assistant, Glasses tab, SSE notifications', 'Built', 'apps/android'),
    ('Android', 'Ray-Ban Meta via DAT', 'Registration, camera stream and photo, capture button, glasses speech', 'Partial', 'glasses/GlassesManager.kt (not run on real hardware)'),
    ('Android', 'Glasses bridge (WebView)', 'Companion app hosting the glasses web app with the DAT bridge', 'Partial', 'apps/glasses-bridge-android'),
    ('Android', 'Inspections step entry, offline outbox, MFA', 'Not in the native app yet', 'Planned', 'apps/android'),
    ('Docs', 'Product documentation', '13-document suite, 99-page Product Book, 60-page PRD, Word + PDF generator', 'Built', 'docs/product'),
]

# ------------------------------------------------------------------ corrections (bugs, gaps, risks)
CORR = [
    # id, area, issue, impact, severity, fix, status
    ('C01', 'Android / repo', 'Root .gitignore "data" excluded the Android data/ source package', 'Fresh clone could not compile', 'Critical', 'Negation rule added; files committed', 'Done'),
    ('C02', 'Android', 'Resolve sent no note; API returns note_required', 'Resolve always failed', 'High', 'Note field added; Resolve disabled until filled', 'Done (verify on device)'),
    ('C03', 'Security / API', 'GET /api/attachments/:id checks org only, not site scope', 'Members can download files of sites they cannot see', 'High', 'Check site scope of the owning record before streaming', 'Open'),
    ('C04', 'Security / DB', 'App connects as owner role, which bypasses RLS', 'RLS is not enforced if a query forgets the tenant filter', 'High', 'Separate migration owner and runtime role (NOBYPASSRLS)', 'Open'),
    ('C05', 'Android', 'Release build signed with the debug key', 'Cannot ship to Play; insecure', 'High', 'Release keystore via env/CI secrets; Play App Signing', 'Open'),
    ('C06', 'Glasses bridge', 'WebView grants every permission and does not restrict navigation', 'Other pages get the native bridge and camera/mic', 'High', 'Allow-list origin, restrict navigation, scoped permission grant', 'Open'),
    ('C07', 'Repo', 'Debug APK committed (apps/glasses-bridge-android/hydris-glasses-debug.apk)', 'Binary in repo, stale build', 'Medium', 'Remove from git history going forward; publish via releases', 'Open'),
    ('C08', 'Security / web', 'Helmet CSP disabled', 'No XSS defence-in-depth', 'Medium', 'Enable CSP with nonces for Next.js', 'Open'),
    ('C09', 'Web', 'Permissions-Policy microphone=() blocks browser speech on /glasses', 'Voice fails on display glasses in browser', 'Medium', 'Allow microphone=(self) on /glasses', 'Open'),
    ('C10', 'Security / API', 'trustProxy: true lets clients spoof X-Forwarded-For', 'Rate limits bypassable on container deploys', 'Medium', 'Trust only the platform proxy hop count', 'Open'),
    ('C11', 'API', 'Rate limiter stores counts per serverless instance', 'Limits ineffective on Vercel', 'Medium', 'Shared store (Upstash Redis or Postgres)', 'Open'),
    ('C12', 'Glasses web', 'Device refresh token and data cache in localStorage', 'Token theft via XSS', 'Medium', 'Move to httpOnly cookie or short-lived token + IndexedDB with encryption', 'Open'),
    ('C13', 'Jobs', 'Vercel cron runs daily (0 3 * * *)', 'Stale-data alarms and pairing expiry only daily', 'High', 'Per-minute scheduler (Vercel Pro cron or external) + idempotent jobs', 'Open'),
    ('C14', 'Ops', '/api/health does not check the database', 'False healthy status', 'Low', 'Add DB ping and version info', 'Open'),
    ('C15', 'Build', 'Node version mismatch: engines 22.x vs Dockerfiles 24', 'Inconsistent runtime', 'Low', 'Align on Node 22 LTS everywhere', 'Open'),
    ('C16', 'Email', 'Email not sent; mail_outbox keeps live reset/invite links with no purge', 'Users cannot reset passwords; sensitive links retained', 'High', 'Wire a provider (Resend/SES), mark sent_at, purge after 7 days', 'Open'),
    ('C17', 'Storage', 'Audit page images stored base64 in Postgres', 'DB bloat, slow queries, cost', 'Medium', 'Move to object storage (Vercel Blob/S3) with signed URLs', 'Open'),
    ('C18', 'Data', 'No retention or purge jobs', 'Unbounded growth, DPDP exposure', 'Medium', 'Retention schedule + purge jobs per table', 'Open'),
    ('C19', 'Design', 'HUD text #141714 on Signal Blue is 3.08:1; Android severity chips 3.2 to 3.8:1', 'Fails WCAG AA', 'Medium', 'White text on blue; darken chip fills', 'Open'),
    ('C20', 'Design', 'High-severity colour differs: #FF9F1C console, #FF5A1F site, #F26414 Android/HUD', 'Inconsistent meaning', 'Low', 'Single token #F26414 shared from one source', 'Open'),
    ('C21', 'Android', 'Action buttons ignore the API allowedActions', 'Users see actions they cannot perform', 'Medium', 'Render buttons from allowedActions', 'Open'),
    ('C22', 'Android', 'Home "All sensors" opens Work on the Alarms tab', 'Wrong navigation', 'Low', 'Navigate to the Sensors tab', 'Open'),
    ('C23', 'Android / DAT', 'stopPreview() stops forwarding frames but not the DAT camera stream', 'Battery drain on glasses', 'Medium', 'Stop the stream / close camera on preview stop', 'Open'),
    ('C24', 'Android', 'Registering glasses always assigns the first visible site', 'Wrong site assignment', 'Medium', 'Site picker during registration', 'Open'),
    ('C25', 'Android', 'Native app and bridge derive different hardware IDs', 'One pair of glasses becomes two devices', 'Medium', 'Single shared hardware-ID derivation', 'Open'),
    ('C26', 'Android', 'Device-token refresh has no mutex', 'Refresh races, forced re-pair', 'Medium', 'Mutex like the user-session refresh', 'Open'),
    ('C27', 'Android build', 'settings.gradle.kts lacks the GitHub Packages repo for DAT; DataStore unused', 'Clean builds can fail', 'Medium', 'Declare repo with read:packages token; remove unused dep', 'Open'),
    ('C28', 'Android', 'No MFA flow; MFA users told to use the web', 'Blocks secure accounts on mobile', 'High', 'Implement TOTP step in LoginScreen', 'Open'),
    ('C29', 'Data', 'Missing FKs on device/actor ids; no GIN index for document search; unused processing status', 'Integrity and search performance', 'Low', 'Migration adding FKs, GIN index; drop or use status', 'Open'),
    ('C30', 'Roles', 'Viewer (auditor persona) lacks audit.read', 'Auditors cannot open the audit log', 'Low', 'Product decision, then update roles.ts', 'Decide'),
    ('C31', 'Deploy', 'Migrations run on cold start (vercel-build unused)', 'Slow cold starts, risky concurrent migrations', 'Medium', 'Run migrations in CI/deploy step', 'Open'),
    ('C32', 'Contracts', 'packages/contracts is empty; Android mirrors JSON by hand', 'Contract drift', 'Medium', 'Generate OpenAPI + Kotlin models', 'Open'),
    ('C33', 'Docs', 'PRD says Resolve sends no note; pairing TTL 10 min in vercel.md vs 300 s; per-minute cron in docs', 'Docs drift', 'Low', 'Update PRD.md, docs/operations/vercel.md', 'Open'),
    ('C34', 'Code comment', 'docdetect.ts says 1240x1754 at 150 dpi; code uses 1654x2339 at 200 dpi', 'Misleading', 'Low', 'Fix comment', 'Open'),
]

# ------------------------------------------------------------------ enhancements
ENH = [
    # id, area, enhancement, value, effort (days), priority, phase
    ('E01', 'CI/CD', 'GitHub Actions: typecheck, unit + API tests, Playwright, web build, Android assemble, migration check', 'Every change verified before deploy', 2, 'P0', 'Phase 1'),
    ('E02', 'Observability', 'Sentry (web, API, Android), structured logs, uptime checks, alerting', 'Find and fix production issues fast', 2, 'P0', 'Phase 2'),
    ('E03', 'Email', 'Transactional email provider for invites, resets, alarm escalations', 'Self-service accounts', 1, 'P0', 'Phase 2'),
    ('E04', 'Storage', 'Object storage for audit images and attachments with signed URLs', 'Scale and cost', 2, 'P0', 'Phase 2'),
    ('E05', 'Jobs', 'Reliable per-minute background jobs (stale data, escalation, pairing expiry, purge)', 'Alarms on time', 1, 'P0', 'Phase 2'),
    ('E06', 'Android', 'Inspections step entry on the phone + offline outbox (Room) with replay', 'Field rounds without signal', 3, 'P0', 'Phase 3'),
    ('E07', 'Android', 'MFA (TOTP) sign-in', 'Secure accounts on mobile', 1, 'P0', 'Phase 3'),
    ('E08', 'Android', 'Unit and UI tests (DocDetect, OCR pairing, repository, screens)', 'Regression safety', 2, 'P0', 'Phase 3'),
    ('E09', 'Android', 'Release signing, R8, Play internal testing track', 'Installable by pilot users', 1, 'P0', 'Phase 3'),
    ('E10', 'Glasses', 'Ray-Ban Meta hardware validation GLS-T01 to T15 (camera, capture button, speech, registration)', 'Proves the glasses path', 3, 'P0', 'Phase 4'),
    ('E11', 'AI / OCR', 'OCR benchmark set (120 pages) and accuracy dashboard', 'Measured accuracy for auditors', 2, 'P1', 'Phase 4'),
    ('E12', 'AI / Voice', 'Voice utterance test set, wake-word tuning, latency targets', 'Reliable hands-free use', 1, 'P1', 'Phase 4'),
    ('E13', 'Contracts', 'OpenAPI generated from zod schemas; typed web client and Kotlin models', 'No contract drift', 2, 'P1', 'Phase 2'),
    ('E14', 'Performance', 'Load tests (k6) for ingest, SSE, dashboard; DB indexes', 'Known capacity', 1, 'P1', 'Phase 4'),
    ('E15', 'Security', 'Dependency and secret scanning, CSP, security headers, pen-test checklist', 'Customer security review ready', 2, 'P0', 'Phase 1'),
    ('E16', 'Glasses', 'Photo steps captured with the glasses camera (FR-045)', 'Evidence without the phone', 1, 'P1', 'Phase 3'),
    ('E17', 'Design', 'Single design-token source (web, HUD, Android) and contrast fixes', 'Consistent, accessible UI', 1, 'P1', 'Phase 3'),
    ('E18', 'Backup / DR', 'Point-in-time restore drill, backup runbook, RPO/RTO targets', 'Recoverability', 1, 'P1', 'Phase 4'),
    ('E19', 'Compliance', 'Audit trail export, e-signature on approvals (Planned design)', 'Regulatory record integrity', 2, 'P2', 'After month'),
    ('E20', 'Localisation', 'Hindi and Tamil for glasses and Android (strings catalogue first)', 'Operator adoption', 2, 'P2', 'After month'),
    ('E21', 'Integrations', 'Edge gateway with MQTT / OPC UA ingestion', 'Real plant data without custom work', 10, 'P2', 'After month'),
    ('E22', 'Reports', 'CPCB/SPCB and OCEMS compliance reports', 'Regulatory value', 5, 'P2', 'After month'),
    ('E23', 'Android', 'Certificate pinning, Play Integrity, root detection', 'Mobile hardening', 1, 'P1', 'Phase 3'),
    ('E24', 'Accessibility', 'Haptics, chart descriptions, larger outdoor text', 'Gloved / outdoor use', 1, 'P2', 'Phase 3'),
]

# ------------------------------------------------------------------ 30-day plan (2 to 3 sessions per day)
# (phase, [ (session title, area, refs, guideline steps, definition of done) ... ]) per day
S = lambda t, a, r, g, d: (t, a, r, g, d)
DAYS = [
    ('Phase 1: Foundations and critical fixes', [
        S('Baseline and branch strategy', 'Repo', 'C07, C15', 'Create dev branch + PR rule; remove committed APK; align Node 22; record current test pass', 'main protected; npm test green; APK out of repo'),
        S('CI pipeline v1', 'CI/CD', 'E01', 'GitHub Actions: install, typecheck, domain + API tests, web build; cache npm', 'CI runs on every PR and passes'),
        S('Review and plan', 'All', '-', 'Triage this sheet; assign owners; set up issue board', 'Board with every C and E item'),
    ]),
    ('Phase 1: Foundations and critical fixes', [
        S('Attachment site-scope check', 'API', 'C03', 'Look up owning record, requireSite before streaming; add API test', 'Test proves cross-site download is 404'),
        S('RLS runtime role', 'DB', 'C04', 'Create app role NOBYPASSRLS; migrations use owner; update config and docs', 'Tests pass with runtime role; RLS enforced'),
        S('CI: Android assemble + Playwright', 'CI/CD', 'E01', 'Add Android job (JDK 17, Gradle cache) and Playwright job against local stack', 'All four CI jobs green'),
    ]),
    ('Phase 1: Foundations and critical fixes', [
        S('Security headers and CSP', 'Web', 'C08, C09, E15', 'Enable CSP with nonces; allow microphone on /glasses; verify voice', 'CSP on; glasses voice works in Chrome'),
        S('Proxy trust and shared rate limits', 'API', 'C10, C11', 'Set trusted hop count; move rate-limit store to Redis/Postgres', 'Limits hold across instances (test)'),
    ]),
    ('Phase 1: Foundations and critical fixes', [
        S('Glasses bridge hardening', 'Android bridge', 'C06', 'Allow-list Hydris origin, block navigation, grant only camera/mic for our origin', 'Foreign page gets no bridge; manual test'),
        S('Android release signing', 'Android', 'C05, E09', 'Release keystore from env/CI secret; R8 on; signed AAB in CI artifact', 'Signed release AAB builds in CI'),
        S('Dependency and secret scanning', 'Security', 'E15', 'Dependabot/renovate, npm audit gate, gitleaks in CI', 'Scans run; no high findings open'),
    ]),
    ('Phase 1: Foundations and critical fixes', [
        S('Device token storage (glasses web)', 'Glasses web', 'C12', 'Move refresh token to httpOnly cookie flow or encrypted IndexedDB; short access tokens', 'No tokens in localStorage'),
        S('Android correctness fixes', 'Android', 'C21, C22, C26', 'Render actions from allowedActions; fix Sensors nav; mutex on device refresh', 'Manual test on emulator; no regressions'),
    ]),
    ('Phase 1: Foundations and critical fixes', [
        S('Week 1 regression and demo check', 'QA', '-', 'Full demo walk-through (web, console, glasses, Android); fix blockers', 'Demo works end to end'),
        S('Docs sync', 'Docs', 'C33, C34', 'Update PRD Resolve, pairing TTL, cron cadence, docdetect comment', 'Docs match code'),
    ]),
    ('Phase 1: Foundations and critical fixes', [
        S('Buffer / catch-up', 'All', '-', 'Close any open Phase 1 items', 'Phase 1 exit: CI green, P0 security fixes merged'),
        S('Weekly retro and plan week 2', 'All', '-', 'Retro; update statuses in this sheet', 'Week 2 owners confirmed'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('Reliable background jobs', 'Jobs', 'C13, E05', 'Per-minute scheduler; idempotent job runner with locking; job run log', 'Stale-data alarm fires within 2 min (test)'),
        S('Health and readiness', 'Ops', 'C14', 'DB ping, migration version, build SHA in /api/health', 'Health fails when DB down'),
        S('Migrations in deploy step', 'Deploy', 'C31', 'Run drizzle migrate in CI before promote; remove cold-start migrate', 'Cold start without migration'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('Email provider', 'Email', 'C16, E03', 'Resend/SES adapter behind mail interface; send outbox; sent_at; templates', 'Invite and reset emails arrive'),
        S('Outbox purge and link expiry', 'Email', 'C16, C18', 'Purge sent mail after 7 days; expire tokens', 'Purge job tested'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('Object storage for audit images', 'Storage', 'C17, E04', 'Blob/S3 adapter; write new pages to storage; signed URLs; PDF reads from storage', 'New scans store no base64 in DB'),
        S('Backfill migration', 'Storage', 'C17', 'Script moving existing base64 pages to storage; verify counts', 'All pages migrated; DB size drops'),
        S('Attachments to storage', 'Storage', 'E04', 'Same adapter for attachments', 'Downloads served by signed URL'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('Observability: errors', 'Observability', 'E02', 'Sentry for API, web and Android; release tags; PII scrubbing', 'Errors visible with stack traces'),
        S('Observability: logs and uptime', 'Observability', 'E02', 'Structured request logs with tenant/request IDs; uptime monitor; alert routes', 'Alert reaches on-call channel'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('Data integrity migration', 'DB', 'C29', 'Add missing FKs, GIN index for search, resolve processing status', 'Migration applied; tests pass'),
        S('Retention schedule', 'Data', 'C18', 'Retention per table (telemetry, events, mail, audit pages) as config + purge jobs', 'Purge jobs run in staging'),
        S('Roles decision', 'Product', 'C30', 'Decide auditor access to audit log; update roles.ts and tests', 'Decision recorded'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('OpenAPI generation', 'Contracts', 'C32, E13', 'Generate OpenAPI from zod route schemas; publish /api/openapi.json', 'Spec covers all routes'),
        S('Typed clients', 'Contracts', 'E13', 'Generate TS client for web and Kotlin models for Android; replace hand models gradually', 'Android compiles with generated models'),
    ]),
    ('Phase 2: Backend hardening and data', [
        S('Staging environment', 'Deploy', '-', 'Staging project + DB branch; seed; env matrix documented', 'Staging URL live'),
        S('Week 2 retro', 'All', '-', 'Update sheet; plan week 3', 'Phase 2 exit: jobs, email, storage, observability live'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Android MFA', 'Android', 'C28, E07', 'TOTP step after password; error states; tests', 'MFA user signs in on phone'),
        S('Android inspections: data layer', 'Android', 'E06', 'Room entities for rounds/steps; repository; op IDs for idempotent replay', 'Unit tests for outbox'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Android inspections: screens', 'Android', 'E06', 'Round list, step-by-step entry, checklist, notes, submit', 'Round completed on emulator'),
        S('Offline replay', 'Android', 'E06', 'WorkManager sync, conflict handling, offline banner', 'Airplane-mode round syncs later'),
        S('Photo steps', 'Android / glasses', 'E16', 'Capture with glasses or phone camera into step attachments', 'Photo step saved'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Android tests', 'Android', 'E08', 'JUnit for DocDetect, OCR row pairing, repository; Compose UI tests for login/home', 'Tests in CI'),
        S('DAT fixes', 'Android', 'C23, C24, C25, C27', 'Stop DAT stream; site picker; shared hardware ID; Gradle repo for DAT', 'Clean build from fresh clone'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Design tokens single source', 'Design', 'C20, E17', 'Tokens JSON -> CSS vars, HUD constants, Kotlin Hx', 'One severity palette everywhere'),
        S('Contrast and accessibility', 'Design', 'C19, E24', 'Fix HUD and chip contrast; chart descriptions; text scale', 'All pairs 4.5:1 or better'),
        S('Mobile hardening', 'Android', 'E23', 'Certificate pinning, Play Integrity, backup rules', 'Pinning verified'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Console polish', 'Web', '-', 'Empty/loading/error states audit; pagination on long lists; keyboard access', 'Console UX checklist passes'),
        S('Glasses web polish', 'Glasses web', '-', 'Overlap check at all viewports; offline queue indicator; voice error states', 'No overlaps; checks scripted'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Play internal testing', 'Android', 'E09', 'Play Console app, internal track, privacy policy, data safety form', 'Testers install from Play'),
        S('Week 3 regression', 'QA', '-', 'Full regression on staging with Android build', 'No P0/P1 bugs open'),
    ]),
    ('Phase 3: Mobile and frontend production quality', [
        S('Buffer / catch-up', 'All', '-', 'Close Phase 3 items', 'Phase 3 exit: Android pilot-ready build'),
        S('Week 3 retro', 'All', '-', 'Update sheet; hardware booked for week 4', 'Glasses and test phones ready'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('Ray-Ban Meta setup', 'Glasses', 'E10', 'Meta AI app, Developer Mode, register Hydris Field; GLS-T01 to T05', 'Registration and permissions pass'),
        S('Camera and capture button', 'Glasses', 'E10', 'GLS-T06 to T10: stream, photo, capture button, guided A4 scan', 'Scan from glasses reaches console'),
        S('Speech and voice', 'Glasses', 'E10, E12', 'GLS-T11 to T15: glasses mic, wake word, TTS through glasses', 'Hey Hydris works on glasses'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('OCR benchmark set', 'AI / OCR', 'E11', 'Build 120-page ground truth; script field/decimal/table accuracy', 'Baseline accuracy report'),
        S('OCR improvements', 'AI / OCR', 'E11', 'Tune thresholds, plausibility rules from benchmark failures', 'Silent error rate at target'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('Voice test set', 'AI / Voice', 'E12', '600 utterances incl. Indian English accents; false accept/reject', 'Voice report vs targets'),
        S('Assistant grounding checks', 'AI', '-', 'Zero answers with numbers not from tools; data age stated', 'Automated check in CI'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('Load and performance tests', 'Performance', 'E14', 'k6: ingest, SSE fan-out, dashboard; fix slow queries', 'Capacity documented; p95 targets met'),
        S('Backup and restore drill', 'Ops', 'E18', 'Point-in-time restore to a branch; runbook timing', 'Restore within RTO target'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('Security review', 'Security', 'E15', 'Pen-test checklist (OWASP ASVS L2), fix findings, threat model update', 'No high findings open'),
        S('Privacy and compliance pack', 'Compliance', '-', 'DPDP notice, retention, DPA, subprocessors list (Vercel, Neon, Anthropic)', 'Pack ready for customer review'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('Pilot onboarding dry run', 'Ops', '-', 'Create pilot org, sites, sensors, users, profiles, pair glasses (runbook)', 'Onboarding under 1 day'),
        S('Training material', 'Ops', '-', 'Operator and admin quick guides, 5-minute videos', 'Materials reviewed'),
    ]),
    ('Phase 4: Hardware validation, AI quality, pilot readiness', [
        S('Buffer / catch-up', 'All', '-', 'Close Phase 4 items', 'All P0 items done'),
        S('Week 4 retro', 'All', '-', 'Update sheet; go/no-go checklist draft', 'Checklist ready'),
    ]),
    ('Phase 5: Release', [
        S('Release candidate', 'Release', '-', 'Freeze; tag RC; deploy to staging; full regression', 'RC signed off by QA'),
        S('Production deploy', 'Release', '-', 'Promote to production; migrations; smoke tests; monitoring watch', 'Production healthy for 24 h'),
        S('Go/no-go review', 'Product', '-', 'Review exit criteria (QA plan); decide pilot start', 'Decision recorded'),
    ]),
    ('Phase 5: Release', [
        S('Hypercare', 'Ops', '-', 'Monitor errors, alarms latency, scan verification turnaround; fix fast', 'No Sev1/Sev2 open'),
        S('Month retrospective and next roadmap', 'All', 'E19 to E22', 'Plan e-signature, localisation, edge gateway, compliance reports', 'Next-month plan agreed'),
    ]),
]

SESSION_TIMES = ['S1 09:30-12:00', 'S2 14:00-16:30', 'S3 17:00-18:30']


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'


def body(ws):
    for r, row in enumerate(ws.iter_rows(min_row=2), 2):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.font = Font(name='Calibri', size=10)
            c.border = BOX
            if r % 2 == 0:
                c.fill = PatternFill('solid', fgColor='F6F8FE')
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()

# Summary
ws = wb.active
ws.title = 'Summary'
ws['A1'] = 'Hydris Field Glasses: Production Readiness Plan'
ws['A1'].font = Font(bold=True, size=16, color=NAVY, name='Arial')
ws['A2'] = 'Inventory of what is built, corrections, enhancements, and a 30-day plan with 2 to 3 work sessions per day.'
ws['A2'].font = Font(italic=True, color='5B6472')
start = dt.date(2026, 10, 27)
rows = [
    ('Plan start', start), ('Plan end', start + dt.timedelta(days=len(DAYS) - 1)), ('Days', len(DAYS)),
    ('Work sessions', '=COUNTA(\'30-Day Plan\'!E2:E200)'),
    ('Sessions done', '=COUNTIF(\'30-Day Plan\'!K2:K200,"Done")'),
    ('Progress', '=IFERROR(B8/B7,0)'),
    ('Backend components', '=COUNTA(Backend!B2:B100)'), ('Frontend components', '=COUNTA(Frontend!B2:B100)'),
    ('Corrections', '=COUNTA(Corrections!A2:A100)'), ('Corrections open', '=COUNTIF(Corrections!G2:G100,"Open")'),
    ('Enhancements', '=COUNTA(Enhancements!A2:A100)'), ('P0 enhancements', '=COUNTIF(Enhancements!F2:F100,"P0")'),
]
for i, (k, v) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True, color=NAVY)
    c = ws.cell(row=i, column=2, value=v)
    if isinstance(v, dt.date): c.number_format = 'dd mmm yyyy'
ws['B9'].number_format = '0%'
ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 18
ws['D4'] = 'Phases'
ws['D4'].font = Font(bold=True, color=NAVY, size=12)
ph = [('Phase 1: Foundations and critical fixes', 'Days 1-7', 'CI, security fixes, Android release signing, bridge hardening'),
      ('Phase 2: Backend hardening and data', 'Days 8-14', 'Jobs, email, object storage, observability, contracts, staging'),
      ('Phase 3: Mobile and frontend production quality', 'Days 15-21', 'Android MFA, inspections + offline, tests, design tokens, Play internal'),
      ('Phase 4: Hardware validation, AI quality, pilot readiness', 'Days 22-28', 'Ray-Ban tests, OCR and voice benchmarks, load, DR, security review'),
      ('Phase 5: Release', 'Days 29-30', 'Release candidate, production, go/no-go, hypercare')]
for i, (a, b, c) in enumerate(ph, 5):
    ws.cell(row=i, column=4, value=a).font = Font(bold=True)
    ws.cell(row=i, column=5, value=b)
    ws.cell(row=i, column=6, value=c)
ws.column_dimensions['D'].width = 50
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 70
ws['D11'] = 'Session guideline'
ws['D11'].font = Font(bold=True, color=NAVY, size=12)
guide = ['S1 (2.5 h, deep work): the main feature or fix of the day.',
         'S2 (2.5 h): second item, tests, or finishing S1.',
         'S3 (1.5 h, optional): review, docs, regression, planning.',
         'Every session: branch, small PR, tests added, CI green, update the Status column.',
         'Definition of done is in each row; do not mark Done without it.']
for i, g in enumerate(guide, 12):
    ws.cell(row=i, column=4, value=g)

# Backend / Frontend
for name, data in (('Backend', BACKEND), ('Frontend', FRONTEND)):
    w = wb.create_sheet(name)
    header(w, ['Area', 'Component', 'What it does', 'Status', 'Location'], [18, 30, 70, 11, 45])
    for r in data: w.append(list(r))
    body(w)

w = wb.create_sheet('Corrections')
header(w, ['ID', 'Area', 'Issue', 'Impact', 'Severity', 'Fix', 'Status'], [6, 16, 55, 38, 10, 50, 16])
for r in CORR: w.append(list(r))
body(w)

w = wb.create_sheet('Enhancements')
header(w, ['ID', 'Area', 'Enhancement', 'Value', 'Effort (days)', 'Priority', 'Phase'], [6, 15, 60, 34, 12, 9, 14])
for r in ENH: w.append(list(r))
body(w)

w = wb.create_sheet('30-Day Plan')
header(w, ['Day', 'Date', 'Weekday', 'Phase', 'Session', 'Task', 'Area', 'Refs', 'Guideline (steps)', 'Definition of done', 'Status', 'Owner', 'Notes'],
       [5, 11, 9, 30, 15, 30, 14, 12, 55, 38, 12, 12, 25])
for d, (phase, sessions) in enumerate(DAYS):
    date = start + dt.timedelta(days=d)
    for i, (t, a, r, g, done) in enumerate(sessions):
        w.append([d + 1, date, date.strftime('%a'), phase, SESSION_TIMES[i], t, a, r, g, done, 'Not started', '', ''])
for c in w['B'][1:]: c.number_format = 'dd mmm'
body(w)
dv = DataValidation(type='list', formula1='"Not started,In progress,Blocked,Done"', allow_blank=True)
w.add_data_validation(dv)
dv.add(f'K2:K{w.max_row}')
for val, col in (('Done', 'C6EFCE'), ('In progress', 'FFEB9C'), ('Blocked', 'FFC7CE')):
    w.conditional_formatting.add(f'K2:K{w.max_row}', CellIsRule(operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=col)))

for sh in ('Corrections',):
    dv2 = DataValidation(type='list', formula1='"Open,In progress,Done,Decide,Done (verify on device)"', allow_blank=True)
    wb[sh].add_data_validation(dv2)
    dv2.add(f'G2:G{wb[sh].max_row}')

wb.save(OUT)
print(OUT, sum(len(s) for _, s in DAYS), 'sessions')
